"""
Table sources.

A source returns ETABS tables as pandas DataFrames:
- column names are ETABS field keys (e.g. "AnalysisSect", "AMod", "UniqueName"),
- identifier columns are str, other columns numeric when possible,
- df.attrs["table"] is the table name, df.attrs["units"] maps column -> unit string in the source's units.

Unit conversion is done by the Engine, not by sources.
"""
import json
import re
from contextlib import contextmanager
from functools import lru_cache
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Tuple

import pandas as pd

from .errors import EngineError, ResultNotAvailable, SourceError, TableNotFound
from .units import parse_unit

ID_COLUMNS = {
    "Story", "Label", "UniqueName", "Name", "Beam", "Column", "Brace", "Pier", "Spandrel", "Element", "Joint",
    "Point", "PointBay", "BeamBay", "ColumnBay", "BraceBay", "WallBay", "FloorBay", "UniquePtI", "UniquePtJ",
    "OutputCase", "CaseType", "StepType", "StepLabel", "AnalysisSect", "DesignSect", "SectProp", "Material",
    "Tower", "SimilarTo", "BSName", "Diaphragm", "PierName", "SpandName", "SpandStory", "Shape", "Type",
    "PropType", "Grade", "GUID", "Color", "Notes", "Location", "LoadPat", "LoadPattern", "ShellObject",
    "ShellElement", "AutoSelect", "DesignType", "Section", "Group", "GroupName",
}


def _is_id(column: str) -> bool:
    return column in ID_COLUMNS or column.startswith("UniquePt")


@lru_cache(maxsize=1)
def _schema() -> Dict[str, Dict[str, str]]:
    path = Path(__file__).with_name("schema.json")
    return json.loads(path.read_text(encoding="utf-8"))["tables"] if path.exists() else {}


def header_to_key(table: str, header: str) -> str:
    """Map an Excel display header to the ETABS field key."""
    header = str(header).strip()
    return _schema().get(table, {}).get(header) or re.sub(r"[\s?]", "", header)


def coerce_types(df: pd.DataFrame) -> pd.DataFrame:
    """Identifier columns -> str; other columns -> numeric when every non-null value parses."""
    for column in df.columns:
        values = df[column]
        if _is_id(column):
            df[column] = values.map(lambda v: v if v is None or (isinstance(v, float) and pd.isna(v)) else
                                    (str(int(v)) if isinstance(v, float) and v.is_integer() else str(v)))
        elif values.dtype == object:
            converted = pd.to_numeric(values, errors="coerce")
            if converted.notna().sum() == values.notna().sum():
                df[column] = converted
    return df


def _with_attrs(df: pd.DataFrame, table: str, units: Dict[str, str]) -> pd.DataFrame:
    df.attrs = {"table": table, "units": {c: u for c, u in units.items() if u and c in df.columns}}
    return df


def _copy(df: pd.DataFrame) -> pd.DataFrame:
    out = df.copy()
    out.attrs = {"table": df.attrs["table"], "units": dict(df.attrs["units"])}
    return out


def _names(value) -> List[str]:
    """Case/combo names as exact strings (no stripping: ETABS names may contain repeated spaces, %, ~, /, ...)."""
    if value is None:
        return []
    return [value] if isinstance(value, str) else [str(v) for v in value]


def _filter_cases(df: pd.DataFrame, cases, combos) -> pd.DataFrame:
    names = _names(cases) + _names(combos)
    if not names or "OutputCase" not in df.columns:
        return df
    out = df[df["OutputCase"].isin(names)].reset_index(drop=True)
    out.attrs = df.attrs
    return out


class TableSource:
    """Base class. Subclasses implement tables() and table()."""
    kind = "base"

    def tables(self) -> List[str]:
        raise NotImplementedError

    def table(self, name: str, cases: Optional[Iterable[str]] = None,
              combos: Optional[Iterable[str]] = None) -> pd.DataFrame:
        raise NotImplementedError

    def close(self) -> None:
        """Release resources (open files). Safe to call more than once."""


class ExcelSource(TableSource):
    """
    One or more workbooks exported by ETABS (Export > Tables to Excel). Sheets are parsed lazily and cached.

    A table present in several files: results tables (with an OutputCase column) are concatenated, converting
    later files to the units of the first; other tables come from the first file that has them.
    """
    kind = "excel"

    def __init__(self, paths):
        import openpyxl
        path_list = [paths] if isinstance(paths, (str, Path)) else list(paths)
        if not path_list:
            raise SourceError("ExcelSource needs at least one file")
        self.paths = [Path(p) for p in path_list]
        self.path = self.paths[0]
        self._books = []
        self._index: Dict[str, List[Tuple[int, str]]] = {}
        for i, path in enumerate(self.paths):
            if not path.exists():
                self.close()
                raise SourceError(f"Excel file not found: {path}")
            try:
                book = openpyxl.load_workbook(path, read_only=True, data_only=True)
            except Exception as e:
                self.close()
                raise SourceError(f"Cannot open Excel file {path}: {e}")
            self._books.append(book)
            for ws in book.worksheets:
                first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), ())
                title = first[0] if first else None
                if isinstance(title, str) and title.strip().upper().startswith("TABLE:"):
                    self._index.setdefault(title.split(":", 1)[1].strip(), []).append((i, ws.title))
        self._cache: Dict[str, pd.DataFrame] = {}

    def tables(self) -> List[str]:
        return list(self._index)

    def table_files(self) -> Dict[str, List[str]]:
        """Table name -> files containing it (in the given order)."""
        return {name: [str(self.paths[i]) for i, _ in entries] for name, entries in self._index.items()}

    def close(self) -> None:
        for book in self._books:
            book.close()

    def table(self, name, cases=None, combos=None):
        if name not in self._index:
            raise TableNotFound(name, "excel")
        if name not in self._cache:
            parts = [self._parse(name, i, sheet) for i, sheet in self._index[name]]
            if len(parts) > 1 and "OutputCase" in parts[0].columns:
                self._cache[name] = _concat_results(name, parts)
            else:
                self._cache[name] = parts[0]
        return _filter_cases(_copy(self._cache[name]), cases, combos)

    def _parse(self, name: str, book: int, sheet: str) -> pd.DataFrame:
        rows = self._books[book][sheet].iter_rows(values_only=True)
        next(rows)
        headers = list(next(rows, ()))
        while headers and headers[-1] is None:
            headers.pop()
        units = list(next(rows, ())) + [None] * len(headers)
        width = len(headers)
        data = [list(r[:width]) + [None] * (width - len(r)) for r in rows if any(v is not None for v in r)]
        keys = [header_to_key(name, h) for h in headers]
        df = coerce_types(pd.DataFrame(data, columns=keys))
        return _with_attrs(df, name, dict(zip(keys, units[:width])))


def _concat_results(name: str, parts: List[pd.DataFrame]) -> pd.DataFrame:
    """Concatenate the same results table from several files, in the units of the first file."""
    units = dict(parts[0].attrs["units"])
    aligned = [parts[0]]
    for part in parts[1:]:
        part = part.copy()
        for column, unit in part.attrs["units"].items():
            target = units.setdefault(column, unit)
            if unit == target:
                continue
            a, b = parse_unit(unit, name, column), parse_unit(target, name, column)
            if a is None or b is None or a[1:] != b[1:]:
                raise SourceError(f"Table '{name}' column '{column}' has incompatible units across files: "
                                  f"'{target}' and '{unit}'")
            part[column] = pd.to_numeric(part[column], errors="coerce") * a[0] / b[0]
        aligned.append(part)
    return _with_attrs(pd.concat(aligned, ignore_index=True), name, units)


def _file_name(table: str) -> str:
    return re.sub(r'[<>:"/\\|?*]', "_", table)


def write_snapshot(folder, frames: Dict[str, pd.DataFrame], units_system: Dict[str, str]) -> List[Path]:
    """Write tables as CSV plus index.json (table -> file, units) for FileSource."""
    folder = Path(folder)
    folder.mkdir(parents=True, exist_ok=True)
    index = {"units_system": units_system, "tables": {}}
    paths = []
    for name, df in frames.items():
        path = folder / f"{_file_name(name)}.csv"
        df.to_csv(path, index=False, encoding="utf-8")
        index["tables"][name] = {"file": path.name, "units": dict(df.attrs.get("units", {}))}
        paths.append(path)
    (folder / "index.json").write_text(json.dumps(index, indent=1, ensure_ascii=False), encoding="utf-8")
    return paths


class FileSource(TableSource):
    """Folder written by Engine.export: one CSV/JSON file per table plus index.json."""
    kind = "file"

    def __init__(self, folder):
        self.folder = Path(folder)
        index = self.folder / "index.json"
        if not index.exists():
            raise SourceError(f"index.json not found in {self.folder}")
        self._index = json.loads(index.read_text(encoding="utf-8"))["tables"]

    def tables(self) -> List[str]:
        return list(self._index)

    def table(self, name, cases=None, combos=None):
        if name not in self._index:
            raise TableNotFound(name, "file")
        entry = self._index[name]
        path = self.folder / entry["file"]
        if path.suffix == ".json":
            df = pd.DataFrame(json.loads(path.read_text(encoding="utf-8")), dtype=object)
        else:
            df = pd.read_csv(path, dtype=str, keep_default_na=True)
        df = coerce_types(df.astype(object).where(df.notna(), None))
        return _filter_cases(_with_attrs(df, name, entry.get("units", {})), cases, combos)


N_MM = 9  # eUnits.N_mm_C


class LiveSource(TableSource):
    """Running ETABS instance (COM). Tables are read in N-mm to keep precision, then converted by the Engine."""
    kind = "live"

    def __init__(self, pid: Optional[int] = None, sap_model=None):
        from .bridge import EtabsDataBridge
        from .connection import get_active_etabs
        self.sap_model = sap_model if sap_model is not None else get_active_etabs(pid=pid)
        self.bridge = EtabsDataBridge(self.sap_model)

    @contextmanager
    def present_units(self, code: int = N_MM):
        previous = self.sap_model.GetPresentUnits()
        if previous != code:
            self.sap_model.SetPresentUnits(code)
        try:
            yield
        finally:
            if previous != code:
                self.sap_model.SetPresentUnits(previous)

    def tables(self) -> List[str]:
        ret = self.sap_model.DatabaseTables.GetAvailableTables()
        return list(ret[1] or [])

    def _select_output(self, cases, combos):
        db = self.sap_model.DatabaseTables
        cases, combos = _names(cases), _names(combos)
        known_cases = set(self.sap_model.LoadCases.GetNameList()[1] or [])
        known_combos = set(self.sap_model.RespCombo.GetNameList()[1] or [])
        unknown = [c for c in cases if c not in known_cases] + [c for c in combos if c not in known_combos]
        if unknown:
            raise EngineError(f"Unknown load cases/combinations: {unknown}")
        db.SetLoadCasesSelectedForDisplay(cases)
        db.SetLoadCombinationsSelectedForDisplay(combos)

    def table(self, name, cases=None, combos=None):
        db = self.sap_model.DatabaseTables
        with self.present_units(N_MM):
            fields = db.GetAllFieldsInTable(name)
            if fields[-1] != 0:
                raise TableNotFound(name, "live")
            keys = list(fields[2])
            if cases is not None or combos is not None:
                self._select_output(cases, combos)
            ret = db.GetTableForDisplayArray(name, [], "", 0, [], 0, [])
        is_result = "OutputCase" in keys
        if ret[-1] != 0:
            if is_result:
                raise ResultNotAvailable(f"No results for table '{name}'. Run the analysis and check cases/combos.")
            raise TableNotFound(name, "live")
        columns = list(ret[2] or [])
        values = list(ret[4] or [])
        width = len(columns)
        rows = [values[i:i + width] for i in range(0, len(values), width)] if width else []
        if is_result and not rows:
            raise ResultNotAvailable(f"No results for table '{name}'. Run the analysis and check cases/combos.")
        df = coerce_types(pd.DataFrame(rows, columns=columns))
        return _with_attrs(df, name, dict(zip(keys, fields[5])))
